"""Self-check for deterministic tool logic (no LLM calls involved).

Run directly: uv run python tests/test_tools.py

search_corpus/retrieve_context's corpus search aren't covered here: they
need a real GOOGLE_API_KEY for embeddings, so they're not something a
static check can honestly verify. Run the agent end to end in `adk web` to
validate that path.
"""

from __future__ import annotations

import asyncio
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from tools import packaging
from tools.intake import _parse_csv, _parse_docx, _parse_text, _parse_xlsx, parse_document
from schemas import DraftList, QuestionList


def test_parse_text():
    assert _parse_text(b"hello world") == "hello world"


def test_parse_csv():
    result = _parse_csv(b"a,b,c\n1,2,3\n")
    assert result == "a | b | c\n1 | 2 | 3", result


def test_parse_docx():
    import docx

    document = docx.Document()
    document.add_paragraph("Question 1: do you support SSO?")
    buf = io.BytesIO()
    document.save(buf)
    text = _parse_docx(buf.getvalue())
    assert "do you support SSO" in text, text


def test_parse_xlsx():
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Question", "Section"])
    sheet.append(["Do you support SCIM?", "Identity"])
    buf = io.BytesIO()
    workbook.save(buf)
    text = _parse_xlsx(buf.getvalue())
    assert "Do you support SCIM?" in text, text
    assert "Identity" in text, text


class _FakeActions:
    skip_summarization = None


class _FakePart:
    def __init__(self, text: str):
        self.text = text
        self.inline_data = None


class _FakeUserContent:
    def __init__(self, text: str):
        self.parts = [_FakePart(text)]


class _FakeToolContext:
    def __init__(self, state: dict, user_content=None):
        self.state = dict(state)
        self.actions = _FakeActions()
        self.user_content = user_content

    async def list_artifacts(self):
        return []


def test_assemble_draft():
    questions = [
        {
            "id": "1",
            "question": "Do you support SSO?",
            "section": "Identity",
            "category": "sso_auth",
        },
        {
            "id": "2",
            "question": "What is your carbon offset policy?",
            "section": "Other",
            "category": "other",
        },
    ]
    drafts = [
        {
            "id": "1",
            "answer": "Yes, SAML SSO is supported.",
            "confidence": "high",
            "needs_sme_review": False,
            "sources": ["past_answers/sso.md"],
        },
        {
            "id": "2",
            "answer": "No relevant material found in the corpus.",
            "confidence": "low",
            "needs_sme_review": True,
            "sources": [],
        },
    ]
    ctx = _FakeToolContext(
        {
            "questions": questions,
            "drafts": {"items": drafts},
            "source_filename": "test.md",
        }
    )
    result = asyncio.run(packaging.assemble_draft(ctx))
    assert "error" not in result, result
    markdown = result["markdown"]
    assert "Do you support SSO?" in markdown
    assert "Yes, SAML SSO is supported." in markdown
    assert "**Needs SME review:** yes, route to Presales" in markdown, markdown
    assert ctx.actions.skip_summarization is True
    assert ctx.state["drafts"] == drafts
    assert ctx.state["final_document"] == markdown


def test_assemble_draft_missing_state():
    ctx = _FakeToolContext({})
    result = asyncio.run(packaging.assemble_draft(ctx))
    assert "error" in result
    assert ctx.actions.skip_summarization is True


def test_output_schemas_have_object_roots():
    assert QuestionList.model_json_schema()["type"] == "object"
    assert DraftList.model_json_schema()["type"] == "object"


def test_parse_document_pasted_text():
    ctx = _FakeToolContext({}, user_content=_FakeUserContent("Question 1: SSO?"))
    result = asyncio.run(parse_document(ctx))
    assert "error" not in result, result
    assert result == {"source": "pasted_text", "characters_parsed": 16}, result
    assert ctx.state["raw_text"] == "Question 1: SSO?"
    assert ctx.state["source_filename"] == "pasted_text"
    assert ctx.actions.skip_summarization is True


def test_parse_document_nothing_found():
    ctx = _FakeToolContext({}, user_content=None)
    result = asyncio.run(parse_document(ctx))
    assert "error" in result
    assert ctx.actions.skip_summarization is True


if __name__ == "__main__":
    checks = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for check in checks:
        check()
        print(f"ok: {check.__name__}")
    print(f"{len(checks)} checks passed")
